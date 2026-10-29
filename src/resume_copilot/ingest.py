"""输入摄取：JD、旧简历、补充信息。

支持 txt / md / docx / pdf / json，图片走多模态（需 RESUME_VISION_MODEL）。
PDF 依赖 pymupdf（可选安装），缺失时给出可执行的替代方案而不是堆栈。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".csv"}
DOCX_EXT = {".docx"}
PDF_EXT = {".pdf"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}

JD_HINTS = ("jd", "岗位", "职位", "招聘", "job", "requirement")
RESUME_HINTS = ("简历", "resume", "cv", "履历")


class IngestError(RuntimeError):
    """输入文件无法读取。"""


@dataclass
class Source:
    name: str
    kind: str  # jd | resume | info
    text: str = ""
    image_path: Path | None = None
    bytes_hint: int = 0
    meta: dict = field(default_factory=dict)

    @property
    def is_empty(self) -> bool:
        return not self.text.strip() and self.image_path is None

    def title(self) -> str:
        return self.name


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise IngestError("读取 .docx 需要 python-docx：pip install python-docx") from exc
    doc = Document(str(path))
    chunks: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(dict.fromkeys(cells)))  # 合并单元格会重复，去重
    return "\n".join(chunks)


def _read_pdf(path: Path) -> str:
    try:
        import pymupdf  # type: ignore
    except ImportError:
        try:
            import fitz as pymupdf  # type: ignore
        except ImportError as exc:
            raise IngestError(
                "读取 .pdf 需要 pymupdf：pip install 'resume-copilot[pdf]'；"
                "或自行把 PDF 另存为 .docx / .md 再喂进来。"
            ) from exc
    with pymupdf.open(str(path)) as doc:
        return "\n".join(page.get_text() for page in doc).strip()


def read_any(path: str | Path, kind: str = "info") -> Source:
    path = Path(path)
    if not path.is_file():
        raise IngestError(f"文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix in IMAGE_EXT:
        return Source(name=path.name, kind=kind, image_path=path, bytes_hint=path.stat().st_size)
    if suffix in TEXT_EXT:
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suffix in DOCX_EXT:
        text = _read_docx(path)
    elif suffix in PDF_EXT:
        text = _read_pdf(path)
    else:
        raise IngestError(f"暂不支持 {suffix}，请转成 txt / md / docx / pdf")
    return Source(name=path.name, kind=kind, text=text.strip(), bytes_hint=path.stat().st_size)


def load_sources(paths: list[str | Path], kind: str) -> list[Source]:
    sources = [read_any(p, kind=kind) for p in paths]
    return [s for s in sources if not s.is_empty]


def guess_kind(name: str) -> str:
    low = name.lower()
    if any(h in low for h in JD_HINTS):
        return "jd"
    if any(h in low for h in RESUME_HINTS):
        return "resume"
    return "info"


def collect_dir(directory: str | Path) -> dict[str, list[Source]]:
    """按文件名猜测类型，把一个目录拆成 jd / resume / info 三堆。"""
    directory = Path(directory)
    if not directory.is_dir():
        raise IngestError(f"目录不存在：{directory}")
    buckets: dict[str, list[Source]] = {"jd": [], "resume": [], "info": []}
    for path in sorted(directory.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in (TEXT_EXT | DOCX_EXT | PDF_EXT | IMAGE_EXT):
            continue
        kind = guess_kind(path.name)
        try:
            src = read_any(path, kind=kind)
        except IngestError:
            continue
        if not src.is_empty:
            buckets[kind].append(src)
    return buckets


def join_text(sources: list[Source], limit: int = 24000) -> str:
    """把多份材料拼成带分隔的纯文本，并做长度保护。"""
    parts: list[str] = []
    for src in sources:
        body = src.text.strip()
        if not body:
            continue
        parts.append(f"### 来源：{src.name}\n{body}")
    text = "\n\n".join(parts)
    return text[:limit]
