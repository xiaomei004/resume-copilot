"""数据驱动的 DOCX 渲染器。

输入是一份 `schemas.normalize_resume()` 产出的结构化简历，
输出是 A4 单页 docx。渲染器不认识"销售"或"研发"，只认识
`header` + `sections[]` + 三种条目类型（role / bullet / text）。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .theme import Theme

PAGE_W_CM = 21.0
PAGE_H_CM = 29.7


# ---------------- 低层工具 ----------------
def _set_font(run, theme: Theme, size: float, bold: bool = False, color: str = "", italic: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color or theme.ink)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(key), theme.font_latin)
    fonts.set(qn("w:eastAsia"), theme.font_east)


def _para(p, theme: Theme, *, before=0.0, after=0.0, line=None, left=0.0, first=0.0, keep=False, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = theme.line if line is None else line
    pf.left_indent = Cm(left)
    pf.first_line_indent = Cm(first)
    pf.widow_control = False
    if keep:
        pf.keep_with_next = True
    if align is not None:
        pf.alignment = align
    return p


def _shade(p, fill: str) -> None:
    ppr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _border(p, theme: Theme, edges=("bottom",), size=8, color="", space=1) -> None:
    ppr = p._element.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color or theme.rule)
        pbdr.append(el)


def _right_tab(p, pos_cm: float) -> None:
    p.paragraph_format.tab_stops.add_tab_stop(Cm(pos_cm), WD_TAB_ALIGNMENT.RIGHT)


# ---------------- 结构件 ----------------
def _heading(doc, title: str, en: str, theme: Theme):
    p = doc.add_paragraph()
    _para(p, theme, before=4.6, after=2.6, line=1.0, keep=True)
    _set_font(p.add_run("  " + title), theme, theme.heading_pt, bold=True, color=theme.accent)
    if en:
        _set_font(p.add_run("   " + en), theme, theme.meta_pt - 1.2, color="93A8C1")
    if theme.band:
        _shade(p, theme.accent_light)
        _border(p, theme, edges=("left",), size=18, color=theme.accent, space=4)
    _border(p, theme, edges=("bottom",), size=6, space=2)
    return p


def _text_width_cm(text: str, pt: float) -> float:
    """粗估一行文字宽度：中日韩字符按 1 em，ASCII 按 0.55 em。"""
    em = pt / 72 * 2.54
    total = 0.0
    for ch in text:
        total += em * (1.0 if ord(ch) > 0x2E80 else 0.55)
    return total


def _role_line(doc, item: dict, theme: Theme, content_w: float):
    """组织｜角色 ……… 日期。

    日期用右对齐制表位推到行尾；字号与制表位都留一点余量，
    否则「组织名 + 角色」稍长就会把日期挤到下一行（这条踩过一次坑）。
    """
    date = item.get("date") or ""
    sep = "  ｜  "
    org_pt = theme.body_pt + 0.35
    role_pt = theme.body_pt - 0.35
    date_pt = theme.meta_pt - 0.45

    # 左半部分太长就把字号和分隔符各收一点，保证日期不被挤到下一行
    budget = content_w - _text_width_cm("  " + date, date_pt) - 0.35
    left = _text_width_cm(item["org"], org_pt) + _text_width_cm(sep + (item.get("role") or ""), role_pt)
    if left > budget:
        sep, org_pt, role_pt = " ｜ ", org_pt - 0.4, role_pt - 0.4
        left = _text_width_cm(item["org"], org_pt) + _text_width_cm(sep + (item.get("role") or ""), role_pt)

    p = doc.add_paragraph()
    _para(p, theme, before=0.4, after=0.6, line=1.1, keep=True)
    _right_tab(p, content_w - 0.05)
    _set_font(p.add_run(item["org"]), theme, org_pt, bold=True)
    if item.get("role"):
        _set_font(p.add_run(sep + item["role"]), theme, role_pt, color=theme.grey)
    if date:
        _set_font(p.add_run("\t" + date), theme, date_pt, color=theme.grey)
    if item.get("stack"):
        sp = doc.add_paragraph()
        _para(sp, theme, before=0.4, after=1.0, line=1.05)
        _set_font(sp.add_run(item["stack"]), theme, theme.meta_pt - 0.4, color=theme.grey)
    if item.get("sub"):
        sb = doc.add_paragraph()
        _para(sb, theme, before=0.4, after=1.2, line=1.05, keep=True)
        _set_font(sb.add_run(item["sub"]), theme, theme.meta_pt - 0.2, bold=True, color=theme.accent)
    return p


def _bullet(doc, b: dict, theme: Theme):
    p = doc.add_paragraph()
    _para(p, theme, before=0.0, after=1.5, left=0.5, first=-0.5)
    _set_font(p.add_run("• "), theme, theme.body_pt, bold=True, color=theme.accent)
    if b.get("lead"):
        _set_font(p.add_run(b["lead"]), theme, theme.body_pt, bold=True)
    _set_font(p.add_run(b["text"]), theme, theme.body_pt)
    return p


def _text_line(doc, text: str, theme: Theme):
    p = doc.add_paragraph()
    _para(p, theme, before=0.6, after=0.8)
    _set_font(p.add_run(text), theme, theme.body_pt)
    return p


def _no_borders(table) -> None:
    tbl_pr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tbl_pr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _fix_table(tbl, widths: list[float]) -> None:
    tbl_pr = tbl._element.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    tbl_pr.append(ind)
    for col, w in zip(tbl.columns, widths, strict=False):
        col.width = Cm(w)
    for cell, w in zip(tbl.rows[0].cells, widths, strict=False):
        tc_pr = cell._element.get_or_add_tcPr()
        tc_w = OxmlElement("w:tcW")
        tc_w.set(qn("w:w"), str(int(w * 567)))
        tc_w.set(qn("w:type"), "dxa")
        tc_pr.append(tc_w)


# ---------------- 头部 ----------------
def _header(doc, resume: dict, theme: Theme, photo: Path | None, content_w: float) -> bool:
    header = resume["header"]
    objective = header.get("objective") or {}
    contacts = [c for c in header.get("contacts", []) if c]
    metas = [m for m in header.get("meta", []) if m]

    if photo is not None and photo.is_file():
        tbl = doc.add_table(rows=1, cols=2)
        _no_borders(tbl)
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl.autofit = False
        _fix_table(tbl, [theme.photo_col_cm, content_w - theme.photo_col_cm])
        cell_photo, cell_text = tbl.rows[0].cells
        pp = cell_photo.paragraphs[0]
        _para(pp, theme, before=3, line=1.0)
        pp.add_run().add_picture(str(photo), width=Cm(theme.photo_cm))
        host = cell_text
        first = host.paragraphs[0]
        used_table = True
    else:
        host = doc
        first = doc.add_paragraph()
        used_table = False

    # 姓名
    _para(first, theme, before=0, after=1, line=1.0)
    _set_font(first.add_run(header.get("name", "")), theme, theme.name_pt, bold=True, color=theme.accent)
    if header.get("name_en"):
        _set_font(first.add_run("   " + header["name_en"]), theme, theme.meta_pt - 0.4, color="9AAFC7")

    # 求职意向
    obj = host.add_paragraph()
    _para(obj, theme, before=3.4, after=1.2, line=1.0)
    _set_font(obj.add_run("求职意向 ｜ " + objective.get("role", "")), theme, theme.body_pt + 0.9, bold=True)
    if objective.get("detail"):
        _set_font(obj.add_run(objective["detail"]), theme, theme.body_pt - 0.1, color=theme.accent)
    if objective.get("meta"):
        _set_font(obj.add_run("　｜　" + objective["meta"]), theme, theme.meta_pt, color=theme.grey)

    # 联系方式
    if contacts:
        cp = host.add_paragraph()
        _para(cp, theme, before=1.2, after=0, line=1.0)
        _set_font(cp.add_run("　｜　".join(contacts)), theme, theme.meta_pt, color=theme.grey)

    # 地点 / 到岗等
    for meta in metas:
        mp = host.add_paragraph()
        _para(mp, theme, before=1.2, after=0, line=1.0)
        _set_font(mp.add_run(meta), theme, theme.meta_pt, color=theme.grey)

    rule = doc.add_paragraph()
    _para(rule, theme, before=5.5, after=0, line=1.0)
    _border(rule, theme, edges=("bottom",), size=12, color=theme.accent, space=1)
    return used_table


def build_docx(resume: dict, theme: Theme, out_path: str | Path, photo: str | Path | None = None) -> dict:
    """渲染一份 docx，返回统计信息。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = theme.font_latin
    normal.font.size = Pt(theme.body_pt)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), theme.font_east)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = theme.line

    sec = doc.sections[0]
    sec.page_width = Cm(PAGE_W_CM)
    sec.page_height = Cm(PAGE_H_CM)
    sec.top_margin = Cm(theme.margin_top)
    sec.bottom_margin = Cm(theme.margin_bottom)
    sec.left_margin = Cm(theme.margin_side)
    sec.right_margin = Cm(theme.margin_side)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.5)
    grid = sec._sectPr.find(qn("w:docGrid"))
    if grid is not None:  # 不关掉网格，中文行距会被吸附到 18pt 倍数
        grid.set(qn("w:type"), "none")
        grid.set(qn("w:linePitch"), "0")

    content_w = PAGE_W_CM - theme.margin_side * 2
    photo_path = Path(photo) if photo else None
    used_table = _header(doc, resume, theme, photo_path, content_w)

    sections = resume.get("sections", [])
    bullets = 0
    for section in sections:
        _heading(doc, section["title"], section.get("en", ""), theme)
        for item in section["items"]:
            if item["type"] == "role":
                _role_line(doc, item, theme, content_w)
                for b in item.get("bullets", []):
                    _bullet(doc, b, theme)
                    bullets += 1
            elif item["type"] == "bullet":
                _bullet(doc, item, theme)
                bullets += 1
            else:
                _text_line(doc, item.get("text", ""), theme)

    doc.save(str(out_path))
    return {
        "path": str(out_path),
        "sections": len(sections),
        "bullets": bullets,
        "with_photo": used_table,
        "body_pt": theme.body_pt,
        "line": theme.line,
        "theme": theme.key,
    }
