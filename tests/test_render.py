"""渲染器：能出 docx、能被重新读回、主题可收紧。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from resume_copilot import schemas
from resume_copilot.render import get_theme, render_resume
from resume_copilot.render.docx_builder import build_docx
from resume_copilot.render.theme import THEMES
from tests._support import load_fixtures


class TestDocxBuilder(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.resume = schemas.normalize_resume(load_fixtures()["writer"])

    def test_build_and_reread(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "resume.docx"
            stats = build_docx(self.resume, get_theme("navy"), out)
            self.assertTrue(out.is_file())
            self.assertGreaterEqual(stats["sections"], 4)
            self.assertGreaterEqual(stats["bullets"], 8)

            doc = Document(str(out))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("林一舟", text)
            self.assertIn("309", text)
            self.assertIn("实习经历", text)
            self.assertAlmostEqual(doc.sections[0].page_width.cm, 21.0, places=1)
            self.assertLess(abs(doc.sections[0].page_height.cm - 29.7), 0.05)

    def test_themes_all_render(self):
        for key in THEMES:
            with tempfile.TemporaryDirectory() as tmp:
                out = Path(tmp) / f"{key}.docx"
                build_docx(self.resume, THEMES[key], out)
                self.assertTrue(out.is_file(), key)

    def test_theme_shrink(self):
        theme = get_theme("navy")
        smaller = theme.shrink(0.95)
        self.assertLess(smaller.body_pt, theme.body_pt)
        self.assertLess(smaller.line, theme.line)

    def test_render_without_photo_is_single_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "nophoto.docx"
            stats = build_docx(self.resume, get_theme("mono"), out, photo=None)
            self.assertFalse(stats["with_photo"])
            self.assertEqual(Document(str(out)).tables, [])


class TestRenderResume(unittest.TestCase):
    def test_render_result_shape(self):
        resume = schemas.normalize_resume(load_fixtures()["writer"])
        with tempfile.TemporaryDirectory() as tmp:
            result = render_resume(resume, tmp, get_theme("brick"), stem="resume", want_pdf=False, want_png=False)
            self.assertIsNotNone(result["docx"])
            self.assertIsNone(result["pdf"])
            self.assertGreater(result["bullets"], 0)
            self.assertIn("theme", result)


if __name__ == "__main__":
    unittest.main()
